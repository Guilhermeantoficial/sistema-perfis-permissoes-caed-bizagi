from __future__ import annotations

import csv
import io
import json
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.models import ProcessSpecification
from app.services.process_service import PERMISSION_LABELS, serialize_process


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_text(cell, text: str, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def process_to_json_bytes(process: ProcessSpecification) -> bytes:
    return json.dumps(serialize_process(process), ensure_ascii=False, indent=2).encode("utf-8")


def process_to_csv_bytes(process: ProcessSpecification) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";", lineterminator="\n")
    writer.writerow(
        [
            "Código do processo",
            "Nome do processo",
            "Área",
            "Responsável",
            "Hierarquia",
            "Perfil",
            "Tipo de agente",
            "Resumo da permissão",
            "Permissões estruturadas",
            "Referência",
            "Observações",
        ]
    )
    for profile in process.profiles:
        codes = [permission.permission_code for permission in profile.permissions]
        writer.writerow(
            [
                process.code,
                process.name,
                process.area,
                process.responsible,
                profile.hierarchy,
                profile.profile_name,
                profile.agent_type,
                profile.permission_summary,
                ", ".join(PERMISSION_LABELS.get(code, code) for code in codes),
                profile.source_reference,
                profile.notes,
            ]
        )
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def process_to_docx_bytes(process: ProcessSpecification) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(20)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Especificação de Perfis e Permissões")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"Processo [{process.code}] {process.name}").bold = True
    document.add_paragraph()

    metadata = document.add_table(rows=6, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.style = "Table Grid"
    fields = [
        ("Documento de origem", process.source_document),
        ("Responsável", process.responsible or "Não informado"),
        ("Área", process.area or "Não informada"),
        ("Status", process.status),
        ("Versão", str(process.version)),
        ("Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    for row, (label, value) in zip(metadata.rows, fields, strict=True):
        _set_cell_text(row.cells[0], label, bold=True)
        _shade_cell(row.cells[0], "E7E9EC")
        _set_cell_text(row.cells[1], value)

    document.add_heading("1. Objetivo", level=1)
    document.add_paragraph(
        "Consolidar, validar e aprovar a matriz de perfis e permissões do Card "
        "Ocorrências, garantindo rastreabilidade e preparação para integração com o Bizagi."
    )

    document.add_heading("2. Matriz de perfis e permissões", level=1)
    landscape = document.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = 1
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    landscape.top_margin = Cm(1.3)
    landscape.bottom_margin = Cm(1.3)
    landscape.left_margin = Cm(1.2)
    landscape.right_margin = Cm(1.2)

    headers = ["Hierarquia", "Perfil", "Tipo de agente", *PERMISSION_LABELS.values(), "Resumo"]
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[index], header, bold=True, size=8)
        _shade_cell(table.rows[0].cells[index], "D9EAD3")

    for profile in process.profiles:
        row = table.add_row().cells
        _set_cell_text(row[0], profile.hierarchy, size=8)
        _set_cell_text(row[1], profile.profile_name, size=8)
        _set_cell_text(row[2], profile.agent_type, size=8)
        codes = {permission.permission_code for permission in profile.permissions}
        for offset, code in enumerate(PERMISSION_LABELS, start=3):
            _set_cell_text(row[offset], "X" if code in codes else "", size=8)
            row[offset].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cell_text(row[-1], profile.permission_summary, size=8)

    document.add_heading("3. Regras aplicadas", level=1)
    for text in (
        "Ações de registrar, editar, excluir, monitorar, aprovar ou administrar implicam a permissão de visualizar.",
        "Aprovar implica monitorar; administrar concede todas as ações do catálogo.",
        "Especificações aprovadas permanecem bloqueadas até uma reabertura formal.",
        "Somente versões aprovadas podem ser enviadas ao conector real do Bizagi.",
        "Cada salvamento, aprovação, reabertura e sincronização gera um evento de auditoria.",
    ):
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("4. Histórico de auditoria", level=1)
    audit_table = document.add_table(rows=1, cols=4)
    audit_table.style = "Table Grid"
    for index, header in enumerate(("Data", "Ação", "Responsável", "Detalhes")):
        _set_cell_text(audit_table.rows[0].cells[index], header, bold=True)
        _shade_cell(audit_table.rows[0].cells[index], "E7E9EC")
    events = sorted(process.audit_events, key=lambda event: event.created_at, reverse=True)
    for event in events[:30]:
        cells = audit_table.add_row().cells
        details = json.loads(event.details_json or "{}")
        _set_cell_text(cells[0], event.created_at.strftime("%d/%m/%Y %H:%M"), size=8)
        _set_cell_text(cells[1], event.action, size=8)
        _set_cell_text(cells[2], event.actor, size=8)
        _set_cell_text(cells[3], json.dumps(details, ensure_ascii=False), size=8)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
